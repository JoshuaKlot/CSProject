"""
Minesweeper Batch Runner (Headless)
====================================
Runs the solver N times with no GUI or socket connection.
Produces an Excel spreadsheet with per-game stats and summaries.
 
Usage:
    python minesweeper_batch_headless.py <num_games> [rows] [cols] [mines]
 
Examples:
    python minesweeper_batch_headless.py 100
    python minesweeper_batch_headless.py 50 16 30 99
    python minesweeper_batch_headless.py 200 9 9 10
"""
 
import random
import sys
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
 
# ── Board helpers ─────────────────────────────────────────────────────────────
 
def generate_board(rows, cols, mine_count, safe_r, safe_c):
    safe_cells = {(safe_r + dr, safe_c + dc)
                  for dr in range(-1, 2) for dc in range(-1, 2)
                  if 0 <= safe_r + dr < rows and 0 <= safe_c + dc < cols}
    candidates = [(r, c) for r in range(rows) for c in range(cols)
                  if (r, c) not in safe_cells]
    mine_count = min(mine_count, len(candidates))
    mines = set(random.sample(candidates, mine_count))
    return mines, mine_count
 
 
def number_board(rows, cols, mines):
    board = {}
    for r in range(rows):
        for c in range(cols):
            if (r, c) in mines:
                board[(r, c)] = 'M'
            else:
                count = sum(1 for dr in range(-1, 2) for dc in range(-1, 2)
                            if (r + dr, c + dc) in mines)
                board[(r, c)] = str(count)
    return board
 
 
def get_neighbours(r, c, rows, cols):
    return [(r + dr, c + dc)
            for dr in range(-1, 2) for dc in range(-1, 2)
            if (dr, dc) != (0, 0) and 0 <= r + dr < rows and 0 <= c + dc < cols]
 
 
def flood_reveal(start_r, start_c, actual, revealed, rows, cols):
    stack = [(start_r, start_c)]
    while stack:
        r, c = stack.pop()
        if (r, c) in revealed:
            continue
        revealed.add((r, c))
        if actual[(r, c)] == '0':
            for nr, nc in get_neighbours(r, c, rows, cols):
                if (nr, nc) not in revealed:
                    stack.append((nr, nc))
 
 
# ── Solver (Becerra 2015) ─────────────────────────────────────────────────────
 
def _propagate(constraints):
    safe = set()
    mines = set()
    changed = True
    while changed:
        changed = False
        surviving = []
        for cells, count in constraints:
            if count == 0:
                newly = cells - safe - mines
                if newly:
                    safe.update(newly)
                    changed = True
            elif count == len(cells):
                newly = cells - safe - mines
                if newly:
                    mines.update(newly)
                    changed = True
            else:
                surviving.append((cells, count))
        constraints = surviving
 
        new_constraints = []
        for i, (ca, cnt_a) in enumerate(constraints):
            for j, (cb, cnt_b) in enumerate(constraints):
                if i >= j:
                    continue
                for (sm, sm_c), (bg, bg_c) in [((ca, cnt_a), (cb, cnt_b)),
                                                ((cb, cnt_b), (ca, cnt_a))]:
                    if sm < bg:
                        diff = bg - sm
                        nc = bg_c - sm_c
                        if 0 <= nc <= len(diff):
                            cand = (frozenset(diff), nc)
                            if cand not in constraints and cand not in new_constraints:
                                new_constraints.append(cand)
                                changed = True
        constraints = constraints + new_constraints
 
        updated = []
        for cells, count in constraints:
            unknown = cells - safe - mines
            if not unknown:
                continue
            nc = count - len(cells & mines)
            if 0 <= nc <= len(unknown):
                updated.append((frozenset(unknown), nc))
        constraints = updated
 
    return safe, mines, constraints
 
 
def solve_step(board, rows, cols):
    constraints = []
    for r in range(rows):
        for c in range(cols):
            cell = board.get((r, c))
            if cell is None or not cell.isdigit():
                continue
            number = int(cell)
            neighbours = get_neighbours(r, c, rows, cols)
            unrevealed = set()
            flagged_count = 0
            for nr, nc in neighbours:
                nb = board.get((nr, nc), 'E')
                if nb == 'E':
                    unrevealed.add((nr, nc))
                elif nb == 'F':
                    flagged_count += 1
            remaining = number - flagged_count
            if remaining < 0 or not unrevealed:
                continue
            constraints.append((frozenset(unrevealed), remaining))
 
    safe, mines, constraints = _propagate(constraints)
    return safe, mines, constraints
 
 
def choose_probabilistic(board, constraints, rows, cols, total_mines, flagged, known_mines):
    known_mines = known_mines or set()
    unrevealed = [(r, c) for r in range(rows) for c in range(cols)
                  if board.get((r, c), 'E') == 'E' and (r, c) not in known_mines]
    if not unrevealed:
        return None
 
    mines_remaining = max(total_mines - len(flagged) - len(known_mines), 0)
    constrained = set()
    for cells, _ in constraints:
        constrained.update(cells)
    constrained -= known_mines
 
    unconstrained = [p for p in unrevealed if p not in constrained]
    bg_prob = min(mines_remaining / len(unconstrained), 1.0) if unconstrained else 1.0
 
    risk = {}
    for pos in unrevealed:
        if pos in constrained:
            fracs = [cnt / len(cells) for cells, cnt in constraints if pos in cells]
            risk[pos] = sum(fracs) / len(fracs)
        else:
            risk[pos] = bg_prob
 
    min_risk = min(risk.values())
    candidates = [p for p, v in risk.items() if v == min_risk]
    return random.choice(candidates)
 
 
# ── Single game ───────────────────────────────────────────────────────────────
 
def run_game(rows, cols, mine_count):
    first_r, first_c = rows // 2, cols // 2
    mines, actual_mine_count = generate_board(rows, cols, mine_count, first_r, first_c)
    actual = number_board(rows, cols, mines)
 
    revealed = set()
    flagged  = set()
    prob_moves = 0
 
    def board_view():
        v = {}
        for r in range(rows):
            for c in range(cols):
                pos = (r, c)
                if pos in flagged:
                    v[pos] = 'F'
                elif pos in revealed:
                    v[pos] = actual[pos]
                else:
                    v[pos] = 'E'
        return v
 
    def click(r, c):
        if (r, c) in revealed or (r, c) in flagged:
            return 'cont'
        if (r, c) in mines:
            return 'lost'
        flood_reveal(r, c, actual, revealed, rows, cols)
        if rows * cols - len(revealed) - len(flagged) == actual_mine_count - len(flagged):
            return 'won'
        return 'cont'
 
    result = click(first_r, first_c)
    if result != 'cont':
        return dict(won=(result == 'won'), prob_moves=0,
                    cells_revealed=len(revealed), total_cells=rows * cols,
                    mines=actual_mine_count, rows=rows, cols=cols)
 
    for _ in range(rows * cols * 4):
        v = board_view()
        safe, det_mines, constraints = solve_step(v, rows, cols)
 
        for pos in det_mines:
            flagged.add(pos)
 
        if safe:
            for pos in sorted(safe):
                if pos not in revealed and pos not in flagged:
                    result = click(*pos)
                    if result in ('lost', 'won'):
                        return dict(won=(result == 'won'), prob_moves=prob_moves,
                                    cells_revealed=len(revealed), total_cells=rows * cols,
                                    mines=actual_mine_count, rows=rows, cols=cols)
            continue
 
        v = board_view()
        safe, det_mines, constraints = solve_step(v, rows, cols)
        pos = choose_probabilistic(v, constraints, rows, cols,
                                   actual_mine_count, flagged, det_mines)
        if pos is None:
            break
        prob_moves += 1
        result = click(*pos)
        if result in ('lost', 'won'):
            return dict(won=(result == 'won'), prob_moves=prob_moves,
                        cells_revealed=len(revealed), total_cells=rows * cols,
                        mines=actual_mine_count, rows=rows, cols=cols)
 
    return dict(won=False, prob_moves=prob_moves, cells_revealed=len(revealed),
                total_cells=rows * cols, mines=actual_mine_count, rows=rows, cols=cols)
 
 
# ── Excel output ──────────────────────────────────────────────────────────────
 
HEADER_FILL  = PatternFill('solid', start_color='1F4E79')
ALT_FILL     = PatternFill('solid', start_color='D6E4F0')
SUMMARY_FILL = PatternFill('solid', start_color='E2EFDA')
WIN_FILL     = PatternFill('solid', start_color='C6EFCE')
LOSS_FILL    = PatternFill('solid', start_color='FFC7CE')
THIN   = Side(style='thin', color='AAAAAA')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
 
 
def style_header(cell):
    cell.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = BORDER
 
 
def style_data(cell, alt=False):
    cell.fill = ALT_FILL if alt else PatternFill()
    cell.font = Font(name='Arial', size=10)
    cell.alignment = Alignment(horizontal='center')
    cell.border = BORDER
 
 
def build_excel(results, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Game Results'
 
    headers = ['Game #', 'Board Size', 'Rows', 'Cols', 'Total Cells',
               'Mines', 'Mine Density (%)', 'Result',
               'Cells Revealed', '% Board Revealed', 'Probabilistic Moves']
    for ci, h in enumerate(headers, 1):
        style_header(ws.cell(row=1, column=ci, value=h))
    ws.row_dimensions[1].height = 20
 
    for i, w in enumerate([8, 12, 7, 7, 12, 7, 16, 9, 14, 17, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
 
    for idx, r in enumerate(results, 1):
        row = idx + 1
        alt = (idx % 2 == 0)
        vals = [idx, f"{r['rows']}x{r['cols']}", r['rows'], r['cols'],
                r['total_cells'], r['mines'], f"=F{row}/E{row}",
                'Won' if r['won'] else 'Lost',
                r['cells_revealed'], f"=I{row}/E{row}", r['prob_moves']]
        for ci, val in enumerate(vals, 1):
            style_data(ws.cell(row=row, column=ci, value=val), alt=alt)
        rc = ws.cell(row=row, column=8)
        rc.fill = WIN_FILL if r['won'] else LOSS_FILL
        rc.font = Font(bold=True, name='Arial', size=10,
                       color='375623' if r['won'] else '9C0006')
        ws.cell(row=row, column=7).number_format  = '0.0%'
        ws.cell(row=row, column=10).number_format = '0.0%'
 
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
 
    ws2 = wb.create_sheet('Summary')
    ws2.column_dimensions['A'].width = 34
    ws2.column_dimensions['B'].width = 20
 
    n = len(results)
    wins = sum(1 for r in results if r['won'])
    ds = 2
 
    summary_rows = [
        ('OVERALL STATISTICS', None),
        ('Total Games Played', n),
        ('Games Won', wins),
        ('Games Lost', n - wins),
        ('Win Rate (%)', '=B3/B2'),
        ('', None),
        ('BOARD CONFIGURATION', None),
        ('Board Size', f"{results[0]['rows']}x{results[0]['cols']}"),
        ('Total Cells', results[0]['total_cells']),
        ('Mine Count', results[0]['mines']),
        ('Average Mine Density (%)', f"=AVERAGE('Game Results'!G{ds}:G{n+1})"),
        ('', None),
        ('PROBABILISTIC MOVES', None),
        ('Total Probabilistic Moves',        f"=SUM('Game Results'!K{ds}:K{n+1})"),
        ('Average Probabilistic Moves/Game', f"=AVERAGE('Game Results'!K{ds}:K{n+1})"),
        ('Min Probabilistic Moves',          f"=MIN('Game Results'!K{ds}:K{n+1})"),
        ('Max Probabilistic Moves',          f"=MAX('Game Results'!K{ds}:K{n+1})"),
        ('', None),
        ('BOARD COVERAGE', None),
        ('Average % Board Revealed', f"=AVERAGE('Game Results'!J{ds}:J{n+1})"),
        ('Min % Board Revealed',     f"=MIN('Game Results'!J{ds}:J{n+1})"),
        ('Max % Board Revealed',     f"=MAX('Game Results'!J{ds}:J{n+1})"),
    ]
 
    pct_labels = {'Win Rate (%)', 'Average Mine Density (%)',
                  'Average % Board Revealed', 'Min % Board Revealed', 'Max % Board Revealed'}
 
    for ri, (label, value) in enumerate(summary_rows, 1):
        lc = ws2.cell(row=ri, column=1, value=label)
        vc = ws2.cell(row=ri, column=2, value=value)
        if value is None and label:
            for cell in (lc, vc):
                cell.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
                cell.fill = HEADER_FILL
                cell.border = BORDER
            lc.alignment = Alignment(horizontal='left', vertical='center')
        elif label:
            for cell in (lc, vc):
                cell.fill = SUMMARY_FILL
                cell.border = BORDER
            lc.font = Font(name='Arial', size=10)
            lc.alignment = Alignment(horizontal='left')
            vc.font = Font(bold=True, name='Arial', size=10)
            vc.alignment = Alignment(horizontal='center')
            if label in pct_labels:
                vc.number_format = '0.0%'
            elif label == 'Average Probabilistic Moves/Game':
                vc.number_format = '0.00'
 
    wb.save(out_path)
 
 
# ── Entry point ───────────────────────────────────────────────────────────────
 
def main():
    if len(sys.argv) < 2:
        print("Usage: python minesweeper_batch_headless.py <num_games> [rows] [cols] [mines]")
        sys.exit(1)
 
    num_games = int(sys.argv[1])
    rows  = int(sys.argv[2]) if len(sys.argv) > 2 else 10
    cols  = int(sys.argv[3]) if len(sys.argv) > 3 else 10
    mines = int(sys.argv[4]) if len(sys.argv) > 4 else max(1, (rows * cols) // 6)
 
    print(f"Running {num_games} games on {rows}x{cols} board with {mines} mines...")
    t0 = time.time()
 
    results = []
    for i in range(num_games):
        r = run_game(rows, cols, mines)
        results.append(r)
        status = 'W' if r['won'] else 'L'
        print(f"  Game {i+1:4d}/{num_games}  {status}  prob_moves={r['prob_moves']:3d}  "
              f"revealed={r['cells_revealed']:4d}/{r['total_cells']}", end='\r')
 
    elapsed = time.time() - t0
    wins = sum(1 for r in results if r['won'])
    print(f"\nDone in {elapsed:.1f}s — {wins}/{num_games} won ({100*wins/num_games:.1f}%)")
 
    out = 'minesweeper_results.xlsx'
    build_excel(results, out)
    print(f"Spreadsheet saved to {out}")
 
 
if __name__ == '__main__':
    main()
